from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt

from app.database import DATA_DIR, get_connection, row_to_dict, rows_to_dicts


EXPORT_DIR = DATA_DIR / "exports"


def _add_multiline(document: Document, text: str) -> None:
    if not text.strip():
        document.add_paragraph("无")
        return
    for line in text.splitlines():
        line = line.strip()
        if line:
            document.add_paragraph(line, style=None)


def build_weekly_summary() -> dict:
    conn = get_connection()
    try:
        profile = row_to_dict(conn.execute("SELECT * FROM project_profile WHERE id = 1").fetchone())
        latest_weekly = row_to_dict(
            conn.execute(
                """
                SELECT w.*, d.name AS document_name
                FROM weekly_reports w
                JOIN documents d ON d.id = w.document_id
                ORDER BY COALESCE(w.period_end, '' ) DESC, w.updated_at DESC
                LIMIT 1
                """
            ).fetchone()
        )
        open_risks = rows_to_dicts(
            conn.execute(
                """
                SELECT * FROM risks
                WHERE status != 'closed'
                ORDER BY CASE level WHEN 'high' THEN 1 WHEN 'medium' THEN 2 ELSE 3 END, id DESC
                LIMIT 8
                """
            ).fetchall()
        )
        active_tasks = rows_to_dicts(
            conn.execute(
                """
                SELECT * FROM tasks
                WHERE status != 'completed'
                ORDER BY
                    CASE priority WHEN 'high' THEN 1 WHEN 'medium' THEN 2 ELSE 3 END,
                    COALESCE(due_date, '') ASC,
                    id DESC
                LIMIT 12
                """
            ).fetchall()
        )
        pending_suggestions = conn.execute(
            "SELECT COUNT(*) AS c FROM update_suggestions WHERE status = 'pending'"
        ).fetchone()["c"]
        return {
            "profile": profile,
            "latest_weekly": latest_weekly,
            "open_risks": open_risks,
            "active_tasks": active_tasks,
            "pending_suggestions": pending_suggestions,
            "generated_at": datetime.now().isoformat(timespec="seconds"),
        }
    finally:
        conn.close()


def export_weekly_docx() -> Path:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    summary = build_weekly_summary()
    profile = summary["profile"] or {}
    latest = summary["latest_weekly"] or {}

    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Microsoft YaHei"
    normal_style.font.size = Pt(10.5)

    document.add_heading("国产化OA集成项目周报汇总", level=1)
    document.add_paragraph(f"项目名称：{profile.get('name', '协同办公平台（OA）集成项目')}")
    document.add_paragraph(f"当前阶段：{profile.get('phase', '开发实施')}")
    document.add_paragraph(f"生成时间：{summary['generated_at']}")
    if latest:
        period = f"{latest.get('period_start') or ''} 至 {latest.get('period_end') or ''}".strip(" 至")
        document.add_paragraph(f"参考周报：{latest.get('document_name', '')} {period}")

    document.add_heading("一、本周工作进展", level=2)
    _add_multiline(document, latest.get("progress_text", "") if latest else "")

    document.add_heading("二、下周工作计划", level=2)
    _add_multiline(document, latest.get("next_plan_text", "") if latest else "")

    document.add_heading("三、需协调及解决事项", level=2)
    _add_multiline(document, latest.get("coordination_text", "") if latest else "")

    document.add_heading("四、存在问题或风险", level=2)
    if summary["open_risks"]:
        for risk in summary["open_risks"]:
            document.add_paragraph(
                f"{risk['title']}：{risk.get('description') or ''} {risk.get('mitigation') or ''}".strip(),
            )
    else:
        _add_multiline(document, latest.get("risk_text", "") if latest else "")

    document.add_heading("五、重点跟踪任务", level=2)
    if summary["active_tasks"]:
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        headers = ["任务", "责任人", "状态", "计划完成", "进度"]
        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header
        for task in summary["active_tasks"]:
            row = table.add_row().cells
            row[0].text = task.get("title") or ""
            row[1].text = task.get("owner") or ""
            row[2].text = task.get("status") or ""
            row[3].text = task.get("due_date") or ""
            row[4].text = f"{task.get('progress', 0)}%"
    else:
        document.add_paragraph("暂无未完成任务。")

    document.add_paragraph(f"待确认智能建议：{summary['pending_suggestions']} 条")
    filename = f"国产化OA集成项目周报汇总-{datetime.now().strftime('%Y%m%d-%H%M%S')}.docx"
    output = EXPORT_DIR / filename
    document.save(output)
    return output
