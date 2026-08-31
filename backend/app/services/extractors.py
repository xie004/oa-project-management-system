from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


MAX_TEXT_CHARS = 60000
MAX_PDF_SIZE = 15 * 1024 * 1024
MAX_PDF_PAGES = 12


@dataclass
class ExtractResult:
    text: str
    status: str
    error: str = ""


def compact_text(text: str, limit: int = MAX_TEXT_CHARS) -> str:
    lines = [clean_line(line) for line in text.splitlines()]
    lines = [line for line in lines if is_meaningful_line(line)]
    counts: dict[str, int] = {}
    for line in lines:
        content = re.sub(r"^\[[^\]]+\]\s*", "", line)
        if len(content) <= 50:
            counts[content] = counts.get(content, 0) + 1
    retained: set[str] = set()
    cleaned_lines: list[str] = []
    for line in lines:
        content = re.sub(r"^\[[^\]]+\]\s*", "", line)
        if counts.get(content, 0) >= 3:
            if content in retained:
                continue
            retained.add(content)
        cleaned_lines.append(line)
    text = "\n".join(cleaned_lines)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:limit]


def clean_line(line: str) -> str:
    line = line.replace("\u3000", " ").replace("\xa0", " ")
    line = re.sub(r"[ \t]+", " ", line)
    if re.fullmatch(r"\s*第\s*\d+\s*页\s*共\s*\d+\s*页\s*", line, re.I):
        return ""
    if re.fullmatch(r"\s*(?:共\s*)?\d+\s*页\s*", line, re.I):
        return ""
    line = re.sub(r"(?:\s*\|\s*)+$", "", line)
    line = re.sub(r"^\s*(?:\|\s*)+", "", line)
    line = re.sub(r"\s*\|\s*(?:\|\s*)+", " | ", line)
    return line.strip(" ;；。")


def is_meaningful_line(line: str) -> bool:
    if not line:
        return False
    content = re.sub(r"^\[[^\]]+\]\s*", "", line)
    if re.fullmatch(r"[\d一二三四五六七八九十]+[.、）)]?", content.strip()):
        return False
    content = re.sub(r"[\s|/\\._\-:：;；,，。]+", "", content)
    return bool(re.search(r"[\w\u4e00-\u9fff]", content)) and len(content) >= 2


def clean_table_cells(values: list[str]) -> list[str]:
    cleaned = [clean_line(value) for value in values]
    while cleaned and not cleaned[-1]:
        cleaned.pop()
    return cleaned


def render_table_row(values: list[str]) -> str:
    values = clean_table_cells(values)
    if not values:
        return ""
    if len(values) == 1:
        return values[0]
    return " | ".join(values)


def find_soffice() -> str | None:
    found = shutil.which("soffice.com") or shutil.which("soffice")
    if found:
        return found
    candidates = [
        r"C:\Program Files\LibreOffice\program\soffice.com",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return candidate
    return None


def convert_with_libreoffice(path: Path, target_ext: str) -> Path | None:
    soffice = find_soffice()
    if not soffice:
        return None
    temp_dir = Path(tempfile.mkdtemp(prefix="oa_pm_convert_"))
    try:
        result = subprocess.run(
            [
                soffice,
                "--headless",
                "--convert-to",
                target_ext.lstrip("."),
                "--outdir",
                str(temp_dir),
                str(path),
            ],
            text=True,
            capture_output=True,
            timeout=80,
        )
        if result.returncode != 0:
            return None
        matches = list(temp_dir.glob(f"*{target_ext}"))
        if not matches:
            return None
        stable_path = temp_dir / f"converted{target_ext}"
        matches[0].replace(stable_path)
        return stable_path
    except Exception:
        shutil.rmtree(temp_dir, ignore_errors=True)
        return None


def cleanup_converted(path: Path | None) -> None:
    if not path:
        return
    parent = path.parent
    if parent.name.startswith("oa_pm_convert_"):
        shutil.rmtree(parent, ignore_errors=True)


def extract_docx(path: Path) -> ExtractResult:
    try:
        doc = Document(str(path))
        parts: list[str] = []
        for paragraph_index, para in enumerate(doc.paragraphs, 1):
            line = clean_line(para.text)
            if is_meaningful_line(line):
                parts.append(f"[段落 {paragraph_index}] {line}")
        for table_index, table in enumerate(doc.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                values = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
                rendered = render_table_row(values)
                if is_meaningful_line(rendered):
                    parts.append(f"[表 {table_index}/行 {row_index}] {rendered}")
        return ExtractResult(compact_text("\n".join(parts)), "indexed")
    except Exception as exc:
        return ExtractResult("", "error", str(exc))


def extract_doc(path: Path) -> ExtractResult:
    converted = convert_with_libreoffice(path, ".docx")
    try:
        if converted:
            return extract_docx(converted)
        return ExtractResult("", "unsupported", "未找到 LibreOffice，无法自动转换 .doc 文件。")
    finally:
        cleanup_converted(converted)


def extract_xlsx(path: Path) -> ExtractResult:
    try:
        workbook = load_workbook(str(path), data_only=True, read_only=True)
        parts: list[str] = []
        for sheet in workbook.worksheets[:5]:
            row_count = 0
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                values = ["" if value is None else str(value).strip() for value in row[:18]]
                rendered = render_table_row(values)
                if is_meaningful_line(rendered):
                    parts.append(f"[工作表 {sheet.title}/行 {row_index}] {rendered}")
                    row_count += 1
                if row_count >= 180:
                    break
        return ExtractResult(compact_text("\n".join(parts)), "indexed")
    except Exception as exc:
        return ExtractResult("", "error", str(exc))


def extract_xls(path: Path) -> ExtractResult:
    converted = convert_with_libreoffice(path, ".xlsx")
    try:
        if converted:
            return extract_xlsx(converted)
        return ExtractResult("", "unsupported", "未找到 LibreOffice，无法自动转换 .xls 文件。")
    finally:
        cleanup_converted(converted)


def extract_pdf(path: Path) -> ExtractResult:
    if path.stat().st_size > MAX_PDF_SIZE:
        return ExtractResult("", "metadata_only", "PDF 文件较大，已登记资料台账，暂不抽取正文。")
    try:
        reader = PdfReader(str(path))
        pages = []
        for page_number, page in enumerate(reader.pages[:MAX_PDF_PAGES], 1):
            for line in (page.extract_text() or "").splitlines():
                cleaned = clean_line(line)
                if is_meaningful_line(cleaned):
                    pages.append(f"[第 {page_number} 页] {cleaned}")
        return ExtractResult(compact_text("\n".join(pages)), "indexed")
    except Exception as exc:
        return ExtractResult("", "error", str(exc))


def extract_text(path: Path) -> ExtractResult:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".doc":
        return extract_doc(path)
    if suffix == ".xlsx":
        return extract_xlsx(path)
    if suffix == ".xls":
        return extract_xls(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix in {".txt", ".md"}:
        try:
            return ExtractResult(compact_text(path.read_text(encoding="utf-8", errors="ignore")), "indexed")
        except Exception as exc:
            return ExtractResult("", "error", str(exc))
    return ExtractResult("", "unsupported", f"暂不支持抽取 {suffix or '无后缀'} 文件正文。")


def normalize_date(raw: str) -> str:
    raw = raw.strip()
    match = re.search(r"(\d{4})[年/-]?(\d{1,2})[月/-]?(\d{1,2})", raw)
    if not match:
        return ""
    year, month, day = match.groups()
    return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"


def extract_period(text: str, filename: str = "") -> tuple[str, str]:
    target = f"{filename}\n{text}"
    match = re.search(
        r"(\d{4})[-年]?(\d{1,2})[-月]?(\d{1,2})[日至\s]*(?:至|-|—|~)(\d{4})?[-年]?(\d{1,2})[-月]?(\d{1,2})",
        target,
    )
    if match:
        y1, m1, d1, y2, m2, d2 = match.groups()
        y2 = y2 or y1
        return f"{int(y1):04d}-{int(m1):02d}-{int(d1):02d}", f"{int(y2):04d}-{int(m2):02d}-{int(d2):02d}"

    match = re.search(r"(\d{4})(\d{2})(\d{2})[-至—~](\d{4})?(\d{2})(\d{2})", target)
    if match:
        y1, m1, d1, y2, m2, d2 = match.groups()
        y2 = y2 or y1
        return f"{y1}-{m1}-{d1}", f"{y2}-{m2}-{d2}"
    return "", ""


def section_between(text: str, start: str, stops: list[str]) -> str:
    start_index = text.find(start)
    if start_index < 0:
        return ""
    start_index += len(start)
    end_index = len(text)
    for stop in stops:
        index = text.find(stop, start_index)
        if index >= 0:
            end_index = min(end_index, index)
    return compact_text(text[start_index:end_index], 12000)


def bullet_lines(text: str, limit: int = 8) -> list[str]:
    lines = []
    for line in re.split(r"[\n\r]+", text):
        cleaned = clean_line(line)
        cleaned = re.sub(r"^\[[^\]]+\]\s*", "", cleaned)
        cleaned = re.sub(
            r"^\s*(?:[\d一二三四五六七八九十]+[.、）)]|[（(]?[\d一二三四五六七八九十]+[）)])\s*",
            "",
            cleaned,
        )
        cleaned = clean_line(cleaned)
        if re.fullmatch(r"[|/\\\s._-]+", cleaned or ""):
            continue
        if len(cleaned) >= 4:
            lines.append(cleaned)
    return lines[:limit]


def parse_weekly_report(text: str, filename: str) -> dict[str, str]:
    progress = section_between(text, "二、本周工作进展", ["三、下周工作计划", "四、需协调", "五、存在问题"])
    next_plan = section_between(text, "三、下周工作计划", ["四、需协调", "五、存在问题"])
    coordination = section_between(text, "四、需协调及解决事项", ["五、存在问题", "六、"])
    risks = section_between(text, "五、存在问题或风险", ["六、"])
    period_start, period_end = extract_period(text, filename)
    return {
        "period_start": period_start,
        "period_end": period_end,
        "progress_text": progress,
        "next_plan_text": next_plan,
        "coordination_text": coordination,
        "risk_text": risks,
    }


def parse_meeting(text: str, filename: str) -> dict[str, str]:
    title = ""
    for line in bullet_lines(text, 5):
        if "会议" in line or "项目" in line:
            title = line[:80]
            break
    if not title:
        title = filename.rsplit(".", 1)[0]

    meeting_time = ""
    match = re.search(r"会议时间[:：]\s*([^\n]+)", text)
    if match:
        meeting_time = normalize_date(match.group(1)) or match.group(1).strip()
    if not meeting_time:
        meeting_time = normalize_date(filename)

    location = ""
    match = re.search(r"会议地点[:：]\s*([^\n]+)", text)
    if match:
        location = match.group(1).strip()

    participants = ""
    match = re.search(r"参会人员[:：]\s*([^\n]+)", text)
    if match:
        participants = match.group(1).strip()

    topics = section_between(text, "二、会议议题", ["三、会议议程", "三、", "四、"])
    decisions = section_between(text, "四、决议事项", ["五、", "六、"])
    if not decisions:
        decisions = section_between(text, "达成共识/结论", ["四、", "五、"])

    return {
        "title": title,
        "meeting_time": meeting_time,
        "location": location,
        "participants": participants,
        "topics": topics,
        "decisions": decisions,
        "summary": compact_text(text, 700),
    }
